"""Excel report generator using OpenPyXL."""

import io
from datetime import datetime
from typing import Any, List, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


class ExcelGenerator:
    """Excel report generator."""

    def __init__(self):
        self.header_font = Font(bold=True, color="FFFFFF")
        self.header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        self.header_alignment = Alignment(horizontal="center", vertical="center")
        self.thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

    def _apply_header_style(self, cell):
        """Apply header styling to a cell."""
        cell.font = self.header_font
        cell.fill = self.header_fill
        cell.alignment = self.header_alignment
        cell.border = self.thin_border

    def _apply_cell_style(self, cell):
        """Apply basic styling to a cell."""
        cell.border = self.thin_border
        cell.alignment = Alignment(vertical="center")

    def _auto_adjust_columns(self, ws):
        """Auto-adjust column widths."""
        for column in ws.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)
            for cell in column:
                try:
                    if cell.value:
                        max_length = max(max_length, len(str(cell.value)))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width

    def generate_inventory_report(
        self, items: List, group_by: Optional[str] = None
    ) -> io.BytesIO:
        """Generate inventory Excel report."""
        wb = Workbook()
        ws = wb.active
        ws.title = "Inventory"

        # Title
        ws["A1"] = "Inventory Report"
        ws["A1"].font = Font(bold=True, size=16)
        ws["A2"] = f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
        ws["A2"].font = Font(italic=True)

        # Headers
        headers = ["Type", "Internal Code", "Description", "Status", "Location", "Storage Date"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=4, column=col, value=header)
            self._apply_header_style(cell)

        # Data
        for row, item in enumerate(items, 5):
            ws.cell(row=row, column=1, value=item.item_type)
            ws.cell(row=row, column=2, value=item.internal_code or "")
            ws.cell(row=row, column=3, value=(item.description or "")[:100])
            ws.cell(row=row, column=4, value=item.status)
            ws.cell(row=row, column=5, value=getattr(item.container, "name", "") if item.container else "")
            ws.cell(row=row, column=6, value=str(item.storage_date) if item.storage_date else "")

            for col in range(1, 7):
                self._apply_cell_style(ws.cell(row=row, column=col))

        self._auto_adjust_columns(ws)

        # Summary sheet
        summary_ws = wb.create_sheet("Summary")
        summary_ws["A1"] = "Summary"
        summary_ws["A1"].font = Font(bold=True, size=14)

        summary_data = [
            ("Total Items", len(items)),
            ("Documents", sum(1 for i in items if i.item_type == "DOCUMENT")),
            ("Equipment", sum(1 for i in items if i.item_type == "EQUIPMENT")),
            ("Consumables", sum(1 for i in items if i.item_type == "CONSUMABLE")),
        ]
        for row, (label, value) in enumerate(summary_data, 3):
            summary_ws.cell(row=row, column=1, value=label)
            summary_ws.cell(row=row, column=2, value=value)

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer

    def generate_movements_report(self, movements: List) -> io.BytesIO:
        """Generate movements Excel report."""
        wb = Workbook()
        ws = wb.active
        ws.title = "Movements"

        ws["A1"] = "Movements Report"
        ws["A1"].font = Font(bold=True, size=16)
        ws["A2"] = f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"

        headers = ["Date", "Type", "Item ID", "Reason", "Performed By", "Notes"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=4, column=col, value=header)
            self._apply_header_style(cell)

        for row, mov in enumerate(movements, 5):
            ws.cell(row=row, column=1, value=mov.movement_date.strftime("%Y-%m-%d %H:%M") if mov.movement_date else "")
            ws.cell(row=row, column=2, value=mov.movement_type)
            ws.cell(row=row, column=3, value=str(mov.stored_item_id) if mov.stored_item_id else "")
            ws.cell(row=row, column=4, value=mov.reason or "")
            ws.cell(row=row, column=5, value=str(mov.performed_by) if mov.performed_by else "")
            ws.cell(row=row, column=6, value=mov.notes or "")

            for col in range(1, 7):
                self._apply_cell_style(ws.cell(row=row, column=col))

        self._auto_adjust_columns(ws)

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer

    def generate_access_requests_report(self, requests: List) -> io.BytesIO:
        """Generate access requests Excel report."""
        wb = Workbook()
        ws = wb.active
        ws.title = "Access Requests"

        ws["A1"] = "Access Requests Report"
        ws["A1"].font = Font(bold=True, size=16)
        ws["A2"] = f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"

        headers = ["Request #", "Status", "Urgency", "Type", "Requested At", "Purpose", "Expected Return"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=4, column=col, value=header)
            self._apply_header_style(cell)

        for row, req in enumerate(requests, 5):
            ws.cell(row=row, column=1, value=req.request_number or "")
            ws.cell(row=row, column=2, value=req.status)
            ws.cell(row=row, column=3, value=req.urgency)
            ws.cell(row=row, column=4, value=req.request_type)
            ws.cell(row=row, column=5, value=req.requested_at.strftime("%Y-%m-%d %H:%M") if req.requested_at else "")
            ws.cell(row=row, column=6, value=(req.purpose or "")[:100])
            ws.cell(row=row, column=7, value=str(req.expected_return_date) if req.expected_return_date else "")

            for col in range(1, 8):
                self._apply_cell_style(ws.cell(row=row, column=col))

        self._auto_adjust_columns(ws)

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer

    def generate_audit_trail_report(self, entries: List) -> io.BytesIO:
        """Generate audit trail Excel report."""
        wb = Workbook()
        ws = wb.active
        ws.title = "Audit Trail"

        ws["A1"] = "Audit Trail Report"
        ws["A1"].font = Font(bold=True, size=16)
        ws["A2"] = f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"

        headers = ["Timestamp", "Event Type", "Table", "Record ID", "User", "Action", "IP Address"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=4, column=col, value=header)
            self._apply_header_style(cell)

        for row, entry in enumerate(entries, 5):
            ws.cell(row=row, column=1, value=entry.timestamp.strftime("%Y-%m-%d %H:%M:%S") if entry.timestamp else "")
            ws.cell(row=row, column=2, value=entry.event_type)
            ws.cell(row=row, column=3, value=entry.table_name or "")
            ws.cell(row=row, column=4, value=str(entry.record_id) if entry.record_id else "")
            ws.cell(row=row, column=5, value=entry.username or "")
            ws.cell(row=row, column=6, value=entry.action or "")
            ws.cell(row=row, column=7, value=str(entry.ip_address) if entry.ip_address else "")

            for col in range(1, 8):
                self._apply_cell_style(ws.cell(row=row, column=col))

        self._auto_adjust_columns(ws)

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer
